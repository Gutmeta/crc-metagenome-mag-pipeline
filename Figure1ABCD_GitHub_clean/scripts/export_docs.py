#!/usr/bin/env python3
"""Generate the manuscript legend, results and Word document from source tables."""
from pathlib import Path
import re
import pandas as pd
from docx import Document
from docx.shared import Pt


def add_markdown(document, text):
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            document.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=2)
        else:
            paragraph = document.add_paragraph()
            for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', line):
                if part.startswith('**') and part.endswith('**'):
                    paragraph.add_run(part[2:-2]).bold = True
                elif part.startswith('*') and part.endswith('*'):
                    paragraph.add_run(part[1:-1]).italic = True
                else:
                    paragraph.add_run(part)


def export_docs(destination):
    ROOT = Path(destination)
    data = pd.read_csv(ROOT/'source_data'/'Source_Data_Figure1C.tsv', sep='\t')
    heterogeneity = pd.read_csv(ROOT/'source_data'/'Source_Data_Figure1B.tsv', sep='\t')
    annotated = pd.read_csv(ROOT/'source_data'/'Source_Data_Figure1A_heterogeneity.tsv', sep='\t')
    details = []
    for row in data.itertuples():
        details.append(f'{row.comparison}: observed = {row.observed_shared_species}, expected = '
                       f'{row.permutation_expected:.2f}, observed/expected = '
                       f'{row.observed_expected_ratio:.2f}, *q* = {row.permutation_bh_q:.6f}')
    panel_c = ('**C,** Observed intersections of selected species sets (bars; labels above bars show observed/expected fold enrichment, ×) and mean '
               'intersections expected under 100,000 fixed-size random draws (peach points). Each set '
               'was sampled independently and uniformly without replacement from its own detectable-species '
               'background: CRC/CRA, 1,813 species; IBD, 1,357; IBS, 465. Selected-set sizes remained '
               '300, 300 and 81, respectively. CR denotes the unified CRC/CRA set. The IBS set comprises '
               'the union of species ranked in the Top 50 by random forest feature importance in the '
               'training model of at least one of five stratified cross-validation folds within the '
               '465-species filtered Mars_2020 background. Species-level *q* values and absolute log2 '
               'fold changes were not thresholded for IBS membership. IBS is represented by one cohort. '
               'Asterisks and dark bars indicate BH-adjusted *q* < 0.05 together with an observed count '
               'above the 95th percentile of the null distribution; these indicate enrichment of '
               'set intersections, not significance of individual species. '
               + '; '.join(details) + '. Empirical upper-tail *p* values used the plus-one correction '
               'and were adjusted across all four comparisons by the Benjamini–Hochberg procedure. '
               'The random seed was 20260903.')
    legend_path = ROOT/'docs'/'Figure1_legend.md'
    legend = legend_path.read_text()
    start = legend.index('**C,**'); end = legend.index('\n\n**D,**', start)
    legend = legend[:start] + panel_c + legend[end:]
    legend_path.write_text(legend)
    results = ('# Figure 1 | Main-text results\n\n## Overlap of selected species sets\n\n'
               'The selected CRC/CRA, IBD and Mars_2020 IBS sets contained 300, 300 and 81 species, '
               'respectively. The IBS set was the union of species ranked in the Top 50 by random '
               'forest feature importance in the training model of at least one of five stratified '
               'cross-validation folds within the 465-species filtered Mars_2020 background. '
               'Species-level q values and absolute log2 fold changes were not thresholded for IBS '
               'membership. To account for differences in detectable-species backgrounds, random sets '
               'of these sizes were drawn independently from their respective eligible backgrounds '
               'in 100,000 simulations. CRC/CRA–IBD, IBD–IBS and the three-way intersection met the '
               'prespecified enrichment criterion; CRC/CRA–IBS did not. '
               + '; '.join(details) + '.\n\n'
               'These results indicate overlap enrichment relative to an independent, uniform, fixed-size '
               'selection model within each eligible background. The reported q values describe '
               'set-overlap enrichment, not significance of individual species. They do not establish concordant '
               'association directions or shared causal mechanisms. IBS evidence is restricted to Mars_2020 '
               'and does not constitute cross-cohort IBS replication.\n')
    dispersion = '; '.join(
        f'{r.disease} (n = {r.number_of_cohorts}): median AUC = {r.median_auc:.3f}, '
        f'IQR = {r.interquartile_range:.3f}, SD = {r.standard_deviation_auc:.3f}, '
        f'range = {r.minimum_auc:.3f}–{r.maximum_auc:.3f}' for r in heterogeneity.itertuples())
    results += ('\n## Within-disease cross-cohort AUC dispersion\n\n' + dispersion + '. '
                'IBD showed the largest observed SD and full range. The middle 50% of AUCs had '
                'similar spreads for CRC and IBD, both wider than CRA. These are descriptive '
                'differences, not formal between-disease tests. CRA had lower dispersion but '
                'its median AUC was close to chance; lower dispersion does not imply stronger '
                'predictive performance. Exploratory CI-derived I-squared and tau-squared '
                'estimates are documented in analysis/auc_heterogeneity. Figure 1A annotates '
                'logit-AUC I² as exploratory values; bootstrap percentile intervals only approximate '
                'standard errors and LODO training datasets overlap.\n')
    results += ('\nAnnotated exploratory I²: ' + '; '.join(
        f'{r.disease}, {r.exploratory_i_squared_percent:.1f}%' for r in annotated.itertuples())
        + '. These are not confirmatory tests or formal comparisons between diseases.\n')
    (ROOT/'docs'/'Figure1_results.md').write_text(results)
    document = Document()
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.comments = ''
    document.core_properties.title = 'Figure 1: methods, legend and results'
    document.styles['Normal'].font.name = 'Arial'
    document.styles['Normal'].font.size = Pt(10)
    for i, name in enumerate(['Figure1_methods.md', 'Figure1_legend.md', 'Figure1_results.md']):
        if i:
            document.add_page_break()
        add_markdown(document, (ROOT/'docs'/name).read_text())
    document.save(ROOT/'docs'/'Figure1_methods_and_legend.docx')


if __name__ == '__main__':
    export_docs(Path(__file__).resolve().parents[1])
