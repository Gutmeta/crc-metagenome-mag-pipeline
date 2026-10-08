#!/usr/bin/env python3
"""Reproduce Figure 1A-D and its source tables from the released figure-level inputs."""
from __future__ import annotations

import argparse
import itertools
import math
import os
from pathlib import Path
import shutil
import sys
import tempfile

sys.dont_write_bytecode = True
os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "figure1_mplconfig"))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Rectangle
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.special import logit
from scipy.stats import beta, chi2, norm

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "input"
SOURCE_DATA = ROOT / "reproduced" / "source_data"
RANDOM_SEED = 20260903
N_OVERLAP_PERMUTATIONS = 100_000
DISEASES = ("CRA", "CRC", "IBD")
SETS = ("CRC_CRA", "IBD", "IBS_Mars_2020")
SET_DISPLAY = {"CRC_CRA": "CRC/CRA", "IBD": "IBD", "IBS_Mars_2020": "IBS (Mars_2020)"}
COMPARISON_DISPLAY = {
    "CRC_CRA|IBD": "CRC/CRA–IBD",
    "CRC_CRA|IBS_Mars_2020": "CRC/CRA–IBS",
    "IBD|IBS_Mars_2020": "IBD–IBS",
    "CRC_CRA|IBD|IBS_Mars_2020": "CRC/CRA–IBD–IBS",
}
COLORS = {"CRA": "#315C80", "CRC": "#218B89", "IBD": "#8671AA"}
INK = "#263745"
MUTED = "#667787"
TEXT = "#111820"


def setup_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.sans-serif": ["Arial"],
            "font.size": 6.0,
            "font.weight": "normal",
            "axes.titlesize": 7.2,
            "axes.titleweight": "normal",
            "axes.labelsize": 6.5,
            "axes.labelweight": "normal",
            "axes.edgecolor": "#CBD5E1",
            "axes.linewidth": 0.8,
            "axes.labelcolor": TEXT,
            "xtick.color": TEXT,
            "ytick.color": TEXT,
            "text.color": TEXT,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
        }
    )


def finalize_figure_typography(fig: plt.Figure) -> None:
    """Apply the requested typography to every title, tick, legend and annotation."""
    from matplotlib.font_manager import FontProperties, findfont
    from matplotlib.text import Text

    # Fail rather than silently exporting a substitute font.
    findfont(FontProperties(family="Arial", style="normal", weight="normal"),
             fallback_to_default=False)
    for artist in fig.findobj(match=Text):
        artist.set_fontfamily("Arial")
        artist.set_fontstyle("normal")
        artist.set_fontweight("normal")
        artist.set_color("#000000")
    fig.patch.set_facecolor("white")
    for ax in fig.axes:
        ax.set_facecolor("white")


def bh_fdr(p_values: np.ndarray) -> np.ndarray:
    p = np.asarray(p_values, dtype=float)
    q = np.full_like(p, np.nan, dtype=float)
    mask = np.isfinite(p)
    if not mask.any():
        return q
    valid = p[mask]
    order = np.argsort(valid)
    ranked = valid[order]
    adjusted = ranked * len(ranked) / np.arange(1, len(ranked) + 1)
    adjusted = np.minimum.accumulate(adjusted[::-1])[::-1]
    restored = np.empty_like(adjusted)
    restored[order] = np.clip(adjusted, 0.0, 1.0)
    q[mask] = restored
    return q


def load_inputs() -> tuple[pd.DataFrame, dict[str, set[str]], dict[str, set[str]]]:
    auc = pd.read_csv(INPUT / "panel_a_lodo_auc.tsv", sep="\t")
    required_auc = {
        "disease",
        "project_accession",
        "auc",
        "ci_lower",
        "ci_upper",
        "heldout_cases",
        "heldout_controls",
        "selected_k",
        "bootstrap_iterations",
        "permutation_p_one_sided",
        "permutation_iterations",
    }
    missing = required_auc.difference(auc.columns)
    if missing:
        raise ValueError(f"panel_a_lodo_auc.tsv is missing columns: {sorted(missing)}")
    auc = auc[auc["disease"].isin(DISEASES)].copy()
    auc["disease"] = pd.Categorical(auc["disease"], categories=DISEASES, ordered=True)
    auc = auc.sort_values(
        ["disease", "auc", "project_accession"],
        ascending=[True, False, True],
    ).reset_index(drop=True)
    observed_counts = auc.groupby("disease", observed=True).size().to_dict()
    if observed_counts != {"CRA": 6, "CRC": 9, "IBD": 9}:
        raise ValueError(f"Unexpected LODO cohort counts: {observed_counts}")

    membership = pd.read_csv(INPUT / "species_sets.tsv", sep="\t")
    background = pd.read_csv(INPUT / "species_detection_backgrounds.tsv", sep="\t")
    for name, table in (("species_sets.tsv", membership), ("species_detection_backgrounds.tsv", background)):
        if not {"analysis_set", "species"}.issubset(table.columns):
            raise ValueError(f"{name} must contain analysis_set and species")
    sets = {
        name: set(group["species"].dropna().astype(str))
        for name, group in membership.groupby("analysis_set")
    }
    backgrounds = {
        name: set(group["species"].dropna().astype(str))
        for name, group in background.groupby("analysis_set")
    }
    if set(sets) != set(SETS) or set(backgrounds) != set(SETS):
        raise ValueError("The clean species inputs must contain CRC_CRA, IBD and IBS_Mars_2020")
    expected_sizes = {"CRC_CRA": 300, "IBD": 300, "IBS_Mars_2020": 81}
    actual_sizes = {name: len(sets[name]) for name in SETS}
    if actual_sizes != expected_sizes:
        raise ValueError(f"Unexpected selected-set sizes: {actual_sizes}")
    for name in SETS:
        missing_species = sets[name] - backgrounds[name]
        if missing_species:
            raise ValueError(f"{name} has {len(missing_species)} species outside its detection background")
    return auc, sets, backgrounds


def heterogeneity_summary(auc: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for disease in DISEASES:
        values = auc.loc[auc["disease"].astype(str).eq(disease), "auc"].to_numpy(dtype=float)
        q1, median, q3 = np.percentile(values, [25, 50, 75])
        rows.append(
            {
                "disease": disease,
                "number_of_cohorts": int(values.size),
                "minimum_auc": float(values.min()),
                "first_quartile_auc": float(q1),
                "median_auc": float(median),
                "third_quartile_auc": float(q3),
                "maximum_auc": float(values.max()),
                "interquartile_range": float(q3 - q1),
                "range": float(values.max() - values.min()),
                "mean_auc": float(values.mean()),
                "standard_deviation_auc": float(values.std(ddof=1)),
            }
        )
    return pd.DataFrame(rows)


def pairwise_matrix(sets: dict[str, set[str]]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for row_set in SETS:
        for column_set in SETS:
            rows.append(
                {
                    "row_set": SET_DISPLAY[row_set],
                    "column_set": SET_DISPLAY[column_set],
                    "shared_species": len(sets[row_set] & sets[column_set]),
                }
            )
    return pd.DataFrame(rows)


def write_source_data(
    auc: pd.DataFrame,
    heterogeneity: pd.DataFrame,
    overlap: pd.DataFrame,
    sets: dict[str, set[str]],
) -> None:
    estimates, _ = exploratory_statistics(auc)
    panel_a_heterogeneity = estimates[estimates.effect_scale.eq('logit_auc')].reset_index(drop=True)
    SOURCE_DATA.mkdir(parents=True, exist_ok=True)
    panel_a_heterogeneity.to_csv(SOURCE_DATA / 'Source_Data_Figure1A_heterogeneity.tsv', sep='\t', index=False)
    panel_a = auc.copy()
    panel_a["disease"] = panel_a["disease"].astype(str)
    panel_a.to_csv(SOURCE_DATA / "Source_Data_Figure1A.tsv", sep="\t", index=False)
    heterogeneity.to_csv(SOURCE_DATA / "Source_Data_Figure1B.tsv", sep="\t", index=False)
    panel_c = overlap.copy()
    panel_c.to_csv(SOURCE_DATA / "Source_Data_Figure1C.tsv", sep="\t", index=False)
    panel_d = pairwise_matrix(sets)
    panel_d.to_csv(SOURCE_DATA / "Source_Data_Figure1D.tsv", sep="\t", index=False)
    notes = pd.DataFrame(
        [
            ("Figure 1A", "ROC-AUC and stratified bootstrap 95% confidence intervals for each held-out project."),
            ("Figure 1B", "Equal-weight descriptive minimum, quartiles, median, maximum, mean and sample SD calculated from Figure 1A AUCs. The IQR remains represented by the thick line in B. Disease headings in A display exploratory logit-AUC I-squared estimates, documented in the Figure 1A heterogeneity sheet."),
            ("Figure 1C", "Expected overlap and empirical upper-tail p values from 100,000 fixed-size random draws, independently from each set's own detection background (1813, 1357 and 465 species); BH correction across four comparisons; seed 20260903. Bar heights are observed counts; labels above bars are observed/expected fold enrichment."),
            ("Figure 1D", "Diagonal values are set sizes; off-diagonal values are pairwise shared-species counts."),
            ("Abbreviation", "CR in Figure 1C and the primary Figure 1D denotes the unified CRC/CRA species set."),
            ("IBS species selection", "The 81 selected IBS species are the union of species ranked in the Top 50 by random forest feature importance in the training model of at least one of five stratified cross-validation folds within the 465-species filtered Mars_2020 background. Species-level q values and absolute log2 fold changes were not thresholded for membership. IBS is represented by one cohort and does not constitute cross-cohort IBS replication."),
            ("Monte Carlo uncertainty", "Tail-probability intervals are exact 95% binomial intervals for Monte Carlo uncertainty, not biological or q-value confidence intervals."),
            ("Significance", "An asterisk denotes BH-adjusted q<0.05 and an observed count above the 95th percentile of the permutation null. These q values describe enrichment of set intersections, not significance of individual species."),
        ],
        columns=["item", "description"],
    )
    with pd.ExcelWriter(SOURCE_DATA / "Source_Data_Figure1.xlsx", engine="openpyxl") as writer:
        panel_a.to_excel(writer, sheet_name="Figure 1A", index=False)
        panel_a_heterogeneity.to_excel(writer, sheet_name="Figure 1A heterogeneity", index=False)
        heterogeneity.to_excel(writer, sheet_name="Figure 1B", index=False)
        panel_c.to_excel(writer, sheet_name="Figure 1C", index=False)
        panel_d.to_excel(writer, sheet_name="Figure 1D", index=False)
        notes.to_excel(writer, sheet_name="Notes", index=False)
        writer.book.properties.creator = ""
        writer.book.properties.lastModifiedBy = ""
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0B3C68")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial", size=9)
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000000"
            for column_cells in sheet.columns:
                width = min(
                    36,
                    max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells) + 2,
                )
                sheet.column_dimensions[column_cells[0].column_letter].width = max(10, width)


def compute_overlap(sets, backgrounds, iterations: int, seed: int):
    names = SETS
    universe = sorted(set.union(*(backgrounds[n] for n in names)))
    ids = {name: i for i, name in enumerate(universe)}
    pools = [np.array([ids[s] for s in sorted(backgrounds[n])]) for n in names]
    sizes = [len(sets[n]) for n in names]
    combos = list(itertools.combinations(range(3), 2)) + [(0, 1, 2)]
    observed = np.array([len(set.intersection(*(sets[names[i]] for i in c))) for c in combos])
    null = np.empty((iterations, 4), dtype=np.int16)
    rng = np.random.default_rng(seed)
    membership = np.zeros((3, len(universe)), dtype=bool)
    for b in range(iterations):
        membership.fill(False)
        for i in range(3):
            membership[i, rng.choice(pools[i], size=sizes[i], replace=False)] = True
        ab = membership[0] & membership[1]
        null[b] = (ab.sum(), (membership[0] & membership[2]).sum(),
                   (membership[1] & membership[2]).sum(), (ab & membership[2]).sum())
    rows = []
    for j, combo in enumerate(combos):
        selected_names = [names[i] for i in combo]
        common = len(set.intersection(*(backgrounds[n] for n in selected_names)))
        theory = common * math.prod(len(sets[n]) / len(backgrounds[n]) for n in selected_names)
        mean = float(null[:, j].mean())
        mean_se = float(null[:, j].std(ddof=1) / np.sqrt(iterations))
        exceed = int(np.sum(null[:, j] >= observed[j]))
        p = (exceed + 1) / (iterations + 1)
        # Exact binomial interval for the underlying null upper-tail probability.
        lo = float(beta.ppf(.025, exceed, iterations-exceed+1)) if exceed else 0.
        hi = float(beta.ppf(.975, exceed+1, iterations-exceed)) if exceed < iterations else 1.
        rows.append(dict(
            comparison_key='|'.join(selected_names),
            comparison=COMPARISON_DISPLAY['|'.join(selected_names)],
            number_of_sets=len(combo),
            set_sizes='|'.join(str(len(sets[n])) for n in selected_names),
            own_background_sizes='|'.join(str(len(backgrounds[n])) for n in selected_names),
            common_background_species=common,
            observed_shared_species=int(observed[j]),
            theoretical_expected=theory,
            permutation_expected=mean,
            expected_monte_carlo_se=mean_se,
            permutation_95th_percentile=float(np.quantile(null[:, j], .95)),
            observed_expected_ratio=float(observed[j]/mean),
            observed_theoretical_expected_ratio=float(observed[j]/theory),
            exceedance_count=exceed,
            permutation_iterations=iterations,
            random_seed=seed,
            permutation_p=p,
            tail_probability_mc_95ci_lower=lo,
            tail_probability_mc_95ci_upper=hi,
        ))
    result = pd.DataFrame(rows)
    result['permutation_bh_q'] = bh_fdr(result.permutation_p.to_numpy())
    result['significant_enrichment'] = (result.permutation_bh_q < .05) & (
        result.observed_shared_species > result.permutation_95th_percentile)
    # Analytical mean is an independent check of the sampler's null model.
    assert np.all(np.abs(result.permutation_expected-result.theoretical_expected)
                  < 6*result.expected_monte_carlo_se)
    return result, null


def exploratory_statistics(auc):
    rows = []
    inputs = []
    z = norm.ppf(.975)
    for disease in DISEASES:
        sub = auc[auc.disease.astype(str).eq(disease)]
        for scale in ('logit_auc', 'raw_auc'):
            transform = logit if scale == 'logit_auc' else np.asarray
            y = transform(sub.auc.to_numpy())
            # These percentile bootstrap intervals are NOT Wald intervals;
            # width-derived SEs are deliberately labelled approximate.
            se = (transform(sub.ci_upper.to_numpy())-transform(sub.ci_lower.to_numpy()))/(2*z)
            if not (np.isfinite(y).all() and np.isfinite(se).all() and (se > 0).all()):
                raise ValueError('Invalid AUC values or CI-derived standard errors')
            v = se**2
            w = 1/v
            center = np.sum(w*y)/w.sum()
            q = float(np.sum(w*(y-center)**2))
            df = len(y)-1
            i2 = max(0., (q-df)/q)*100 if q > 0 else 0.
            c = w.sum() - np.sum(w**2)/w.sum()
            tau2 = max(0., (q-df)/c)
            rows.append(dict(disease=disease, number_of_cohorts=len(y), effect_scale=scale,
                             cochran_q=q, degrees_of_freedom=df,
                             exploratory_q_p=float(chi2.sf(q,df)),
                             exploratory_i_squared_percent=i2,
                             exploratory_tau_squared_dl=tau2,
                             exploratory_tau_dl=float(np.sqrt(tau2)),
                             se_method='95% percentile-CI width / (2 * normal 97.5th percentile)',
                             assumption='Independent cohort estimates; approximate normality on chosen scale',
                             interpretation='Exploratory only: shared LODO training and approximate SEs'))
            for accession, effect, standard_error in zip(sub.project_accession,y,se):
                inputs.append(dict(disease=disease, project_accession=accession, effect_scale=scale,
                                   transformed_auc=effect, approximate_se=standard_error))
    return pd.DataFrame(rows), pd.DataFrame(inputs)


def clean(ax, grid='x'):
    ax.spines[['top', 'right', 'left']].set_visible(False)
    ax.spines['bottom'].set_color('#BCC9D3')
    ax.spines['bottom'].set_linewidth(.65)
    ax.tick_params(length=0, pad=2, labelsize=5.5)
    ax.set_axisbelow(True)
    ax.grid(axis=grid, color='#E6ECF0', lw=.55)


def render_figure(auc, sets, summary, overlap, out, stem='Figure1_ABCD'):
    estimates, _ = exploratory_statistics(auc)
    i_squared = estimates[estimates.effect_scale.eq('logit_auc')].set_index('disease').exploratory_i_squared_percent
    setup_style()
    plt.rcParams.update({'font.family': 'Arial', 'font.sans-serif': ['Arial'], 'font.weight': 'normal', 'font.style': 'normal',
                         'font.size': 6, 'text.color': INK, 'axes.labelcolor': INK,
                         'xtick.color': MUTED, 'ytick.color': INK})
    fig = plt.figure(figsize=(170/25.4, 55/25.4))
    fig.text(.017, .927, 'A', weight='normal', fontsize=8)
    grid = fig.add_gridspec(1, 3, left=.036, right=.554, bottom=.19, top=.87, wspace=.16)
    for k, disease in enumerate(DISEASES):
        inner = grid[k].subgridspec(1, 2, width_ratios=[1.32, 1], wspace=.035)
        labels = fig.add_subplot(inner[0]); ax = fig.add_subplot(inner[1])
        sub = auc[auc.disease.astype(str).eq(disease)].reset_index(drop=True)
        for a in (labels, ax):
            a.set_ylim(len(sub)-.6, -.6)
        for y, row in enumerate(sub.itertuples()):
            labels.text(.98, y, row.project_accession, ha='right', va='center', fontsize=5.5)
            ax.errorbar(row.auc, y, xerr=[[row.auc-row.ci_lower], [row.ci_upper-row.auc]],
                        fmt='o', ms=3.3, color=COLORS[disease], ecolor=COLORS[disease],
                        elinewidth=.8, capsize=1.5, markeredgecolor='white', markeredgewidth=.35, zorder=3)
        labels.set_xlim(0, 1); labels.axis('off')
        fig.text(labels.get_position().x0, .927, disease, color=COLORS[disease], fontsize=7.3, weight='normal')
        fig.text(ax.get_position().x1, .927, f'I² = {i_squared[disease]:.1f}%', ha='right', fontsize=5.4, color=MUTED)
        ax.set_xlim(0, 1.04); ax.set_xticks([0, .5, 1], ['0', '0.5', '1.0']); ax.set_yticks([])
        ax.set_xlabel('LODO AUC', fontsize=6, labelpad=3)
        clean(ax); ax.axvline(.5, color='#97A7B5', lw=.7, ls=(0, (3, 3)))
    fig.text(.585, .927, 'B', weight='normal', fontsize=8)
    fig.text(.619, .927, 'LODO AUC heterogeneity', weight='normal', fontsize=7)
    ax = fig.add_axes([.619, .62, .36, .255])
    rng = np.random.default_rng(RANDOM_SEED)
    for y, disease in enumerate(DISEASES):
        row = summary[summary.disease.eq(disease)].iloc[0]
        values = auc.loc[auc.disease.astype(str).eq(disease), 'auc'].to_numpy()
        ax.plot([row.minimum_auc, row.maximum_auc], [y,y], color='#A3B2BF', lw=1.1)
        ax.plot([row.first_quartile_auc, row.third_quartile_auc], [y,y], color=COLORS[disease], lw=4, alpha=.35)
        ax.scatter(values, y+rng.uniform(-.075,.075,len(values)), s=12, color=COLORS[disease], edgecolor='white', linewidth=.5, zorder=3)
        ax.scatter(row.median_auc, y, marker='D', s=22, color=COLORS[disease], edgecolor=INK, linewidth=.6, zorder=4)
    ax.set_ylim(2.45,-.45); ax.set_xlim(.3,1.02)
    ax.set_yticks(range(3), DISEASES); ax.set_xticks([.5,.75,1], ['0.50','0.75','1.00'])
    ax.set_xlabel('Held-out AUC', fontsize=6, labelpad=2)
    clean(ax); ax.axvline(.5, color='#97A7B5', lw=.7, ls=(0,(3,3)))
    fig.text(.585, .422, 'C', weight='normal', fontsize=8)
    fig.text(.619, .422, 'Shared species', weight='normal', fontsize=7)
    ax = fig.add_axes([.619, .16, .195, .215])
    x = np.arange(4)
    ax.bar(x, overlap.observed_shared_species, width=.58,
           color=['#315C80' if v else '#BDD0DE' for v in overlap.significant_enrichment], zorder=3)
    ax.scatter(x, overlap.permutation_expected, s=17, color='#D99168', edgecolor='white', linewidth=.6, zorder=4)
    for i, row in enumerate(overlap.itertuples()):
        ax.text(i, row.observed_shared_species+2, f'{row.observed_expected_ratio:.2f}×'+('*' if row.significant_enrichment else ''),
                ha='center', va='bottom', fontsize=5.5)
    ax.set_ylim(0,85); ax.set_yticks([0,40,80]); ax.set_xticks(x,['CR–IBD','CR–IBS','IBD–IBS','All 3'])
    clean(ax,'y'); ax.tick_params(axis='x', labelsize=4.8)
    fig.legend(handles=[Patch(facecolor='#315C80',label='Observed'),
                        Line2D([],[],marker='o',ls='',color='#D99168',markersize=3,label='Expected')],
               loc='lower left', bbox_to_anchor=(.616,.009), frameon=False, ncol=2,
               fontsize=5, handlelength=.9, columnspacing=.9, borderaxespad=0)
    fig.text(.842, .422, 'D', weight='normal', fontsize=8)
    fig.text(.865, .422, 'Pairwise overlap', weight='normal', fontsize=6.5)
    ax = fig.add_axes([.884,.158,.096,.225])
    for i,left in enumerate(SETS):
        for j,right in enumerate(SETS):
            ax.add_patch(Rectangle((j-.5,i-.5),1,1,facecolor='#EDF1F4' if i==j else '#BBD4DF',edgecolor='white',lw=1))
            ax.text(j,i,str(len(sets[left]&sets[right])),ha='center',va='center',fontsize=5.8)
    ax.set(xlim=(-.5,2.5),ylim=(2.5,-.5),aspect='equal')
    ax.set_xticks(range(3),['CR','IBD','IBS']); ax.set_yticks(range(3),['CR','IBD','IBS'])
    ax.tick_params(length=0,labelsize=5.2,pad=2)
    for spine in ax.spines.values(): spine.set_visible(False)
    fig.text(.843,.065,'IBS: Mars_2020',fontsize=4.8,color=MUTED)
    fig.text(.843,.029,'CR = CRC/CRA',fontsize=4.8,color=MUTED)
    out.mkdir(parents=True, exist_ok=True)
    finalize_figure_typography(fig)
    for ext, kwargs in [('png',{'dpi':300}),('pdf',{}),('svg',{}),('tiff',{'dpi':600,'pil_kwargs':{'compression':'tiff_lzw'}})]:
        fig.savefig(out/f'{stem}.{ext}',**kwargs)
    plt.close(fig)
    print(f'Saved {stem}: 170 × 55 mm.')


def export_analysis(auc, summary, overlap, null_counts, destination):
    overlap_dir = destination / "analysis" / "overlap"
    heterogeneity_dir = destination / "analysis" / "auc_heterogeneity"
    overlap_dir.mkdir(parents=True, exist_ok=True)
    heterogeneity_dir.mkdir(parents=True, exist_ok=True)
    overlap.to_csv(overlap_dir / "Overlap_results.tsv", sep="\t", index=False)
    np.savez_compressed(overlap_dir / "Overlap_null_counts.npz", counts=null_counts,
                        comparisons=overlap.comparison.to_numpy(dtype=str),
                        seed=RANDOM_SEED, iterations=N_OVERLAP_PERMUTATIONS)
    estimates, inputs = exploratory_statistics(auc)
    summary.to_csv(heterogeneity_dir / "AUC_descriptive_heterogeneity.tsv", sep="\t", index=False)
    estimates.to_csv(heterogeneity_dir / "AUC_exploratory_heterogeneity.tsv", sep="\t", index=False)
    inputs.to_csv(heterogeneity_dir / "AUC_approximate_standard_errors.tsv", sep="\t", index=False)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "reproduced",
                        help="Destination for regenerated artifacts (default: reproduced/).")
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    # Keep distributed reference figures and inputs separate from generated outputs.
    if destination == ROOT or destination in ROOT.parents:
        parser.error("Use a separate output directory, not the release root or its parents.")
    reserved = [ROOT / n for n in ("figures", "input", "source_data", "docs", "analysis", "scripts")]
    if any(destination == p or p in destination.parents for p in reserved):
        parser.error("The output directory must not be inside a released reference-data folder.")
    from matplotlib.font_manager import FontProperties, findfont
    findfont(FontProperties(family="Arial", style="normal", weight="normal"), fallback_to_default=False)
    global SOURCE_DATA
    SOURCE_DATA = destination / "source_data"
    auc, sets, backgrounds = load_inputs()
    overlap, null_counts = compute_overlap(sets, backgrounds, N_OVERLAP_PERMUTATIONS, RANDOM_SEED)
    if overlap.observed_shared_species.tolist() != [66, 18, 25, 8]:
        raise ValueError("Unexpected species intersections")
    summary = heterogeneity_summary(auc)
    write_source_data(auc, summary, overlap, sets)
    render_figure(auc, sets, summary, overlap, destination / "figures")
    export_analysis(auc, summary, overlap, null_counts, destination)
    docs = destination / "docs"
    docs.mkdir(parents=True, exist_ok=True)
    for name in ("Figure1_methods.md", "Figure1_legend.md", "Data_dictionary.md", "Data_availability.md"):
        shutil.copyfile(ROOT / "docs" / name, docs / name)
    from export_docs import export_docs
    export_docs(destination)
    print("Reproduced Figure 1, source tables, analysis records and manuscript documents.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
